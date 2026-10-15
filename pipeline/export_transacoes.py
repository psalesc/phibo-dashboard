#!/usr/bin/env python3
"""Exporta as transações de um mês (histórico ou corrente, mesmo shape
usado pelo painel interativo) para uma planilha .xlsx simples, ordenada por
data, pronta pra abrir no Excel/Sheets.

Uso: python3 export_transacoes.py <dados.json> <mes_ano_label> <saida.xlsx>
`dados.json` é um arquivo com o mesmo shape de historical/<mes>.json (tem
que ter a chave "records").
"""
import sys
import json
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = 'Arial'


def sort_key(rec):
    parts = (rec.get('data') or '').split('/')
    if len(parts) == 3:
        try:
            return (int(parts[2]), int(parts[1]), int(parts[0]))
        except ValueError:
            pass
    return (9999, 99, 99)


def main():
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)
    data_path, mes_ano, out_path = sys.argv[1], sys.argv[2], sys.argv[3]

    with open(data_path, encoding='utf-8') as f:
        data = json.load(f)
    records = sorted(data['records'], key=sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Transações'

    headers = ['Vendedor', 'Data', 'Cliente', 'Telefone', 'Categoria', 'Plano/Serviço', 'Valor (R$)', 'Status', 'Observação']
    header_fill = PatternFill('solid', fgColor='1C58A8')
    header_font = Font(name=FONT_NAME, bold=True, color='FFFFFF', size=11)
    thin = Side(style='thin', color='D9D9D9')
    border = Border(bottom=thin)

    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for rec in records:
        row = [
            rec.get('vendedor') or '',
            rec.get('data') or '',
            rec.get('cliente') or '',
            (rec.get('telefone') or '').strip(),
            rec.get('categoria') or '',
            rec.get('plano') or '',
            rec.get('valor'),
            rec.get('status') or '',
            rec.get('observacao') or '',
        ]
        ws.append(row)

    n = ws.max_row
    for r in range(2, n + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = border
            if c == 7:  # Valor
                cell.number_format = '#,##0.00'

    widths = [12, 12, 24, 18, 12, 30, 12, 10, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{n}'

    # rodapé com total e nota de origem
    total = sum(r.get('valor') or 0 for r in records)
    ws.cell(row=n + 2, column=1, value='Total de transações:').font = Font(name=FONT_NAME, bold=True, size=10)
    ws.cell(row=n + 2, column=2, value=len(records)).font = Font(name=FONT_NAME, size=10)
    ws.cell(row=n + 3, column=1, value='Soma dos valores (R$):').font = Font(name=FONT_NAME, bold=True, size=10)
    ws.cell(row=n + 3, column=2, value=round(total, 2)).font = Font(name=FONT_NAME, size=10)
    note = (f'Transação a transação de {mes_ano}, extraído diretamente das planilhas de Samuel, Marina e Bruno '
            f'(a soma dos valores acima não bate com o "Faturamento Total" do painel — o painel usa a Controle '
            f'Vendas/fórmula ×4 pros mensais, esta planilha mostra o valor real de cada venda).')
    ws.cell(row=n + 5, column=1, value=note).font = Font(name=FONT_NAME, italic=True, size=9, color='666666')

    wb.save(out_path)
    print(f'OK: {out_path} gerado com {len(records)} transações, total R$ {total:,.2f}')


if __name__ == '__main__':
    main()
